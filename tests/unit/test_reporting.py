"""Daily report: workbook, Markdown summary, metrics, idempotency and security."""

from __future__ import annotations

import json
import zipfile
from datetime import date

import openpyxl
import pytest

from job_agent.reporting import queries as q
from job_agent.reporting.service import generate_daily_report, list_reports, report_paths, resolve_report_file
from job_agent.reporting.workbook import SHEET_NAMES

D = date(2026, 10, 10)  # report date (America/New_York = UTC-4 in October)
EVIL_TITLE = '=HYPERLINK("http://evil.example/x","click me")'
RAW_SECRET = "Bearer abcdefghijklmnopqrstuvwxyz0123456789"


# --------------------------------------------------------------------------- #
# Seeding helpers (direct SQL matching 001_initial.sql)
# --------------------------------------------------------------------------- #
def _score(db, job_id, score, category, *, shortlisted=0, flags=None, reasons=None, missing=None, current=1,
           version="v1"):
    return db.insert("job_scores", {
        "job_id": job_id, "scoring_version": version, "profile_hash": "p", "config_hash": "c",
        "content_hash": f"h{job_id}{version}", "overall_score": score, "category": category,
        "shortlisted": shortlisted, "factors": "[]", "eligibility_flags": json.dumps(flags or []),
        "reasons_to_apply": json.dumps(reasons or []), "missing_qualifications": json.dumps(missing or []),
        "is_current": current, "scored_at": "2026-10-10T16:00:00+00:00",
    })


def _app(db, job_id, uid, status, *, confirmation="none", submitted_at=None, prepared_at=None,
         created_at="2026-10-01T12:00:00+00:00", notes=None):
    return db.insert("applications", {
        "application_uid": uid, "job_id": job_id, "status": status, "status_updated_at": created_at,
        "confirmation_status": confirmation, "submitted_at": submitted_at, "prepared_at": prepared_at,
        "last_activity_at": created_at, "created_at": created_at, "updated_at": created_at, "notes": notes,
        "next_action": "Review" if status == "awaiting_approval" else None,
    })


def _event(db, app_id, to_status, at, *, evidence=None, actor="system"):
    db.insert("application_events", {
        "application_id": app_id, "event_type": "status_change", "to_status": to_status, "actor": actor,
        "details": "{}", "evidence": json.dumps(evidence) if evidence else None, "created_at": at,
    })


@pytest.fixture
def seeded(db, make_job):
    ids = {}
    old = "2026-10-01T12:00:00+00:00"
    ids["j1"] = make_job(
        employer="Acme Analytics", title="Business Analyst", discovered_at="2026-10-10T15:00:00+00:00",
        pipeline_state="shortlisted", salary_min=80000, salary_max=95000, salary_currency="USD",
        salary_period="year", salary_is_base=1, posted_at="2026-10-08", job_url="https://example.com/jobs/1",
        job_url_canonical="https://example.com/jobs/1", application_url="javascript:alert(1)",
        last_checked_at="2026-10-10T15:00:00+00:00", closing_at="2026-10-13T23:00:00+00:00")
    ids["j2"] = make_job(employer="Globex", title=EVIL_TITLE, discovered_at="2026-10-11T03:30:00+00:00",
                         job_url="https://globex.example/careers/2", job_url_canonical="https://globex.example/2",
                         application_url="https://globex.example/apply/2")
    ids["j3"] = make_job(employer="Initech", title="Data Analyst", discovered_at="2026-10-10T02:00:00+00:00",
                         pipeline_state="scored", job_url="ftp://initech.example/job3", job_url_canonical=None)
    ids["j4"] = make_job(employer="Late Corp", title="Product Analyst", discovered_at="2026-10-11T05:00:00+00:00")
    ids["j5"] = make_job(employer="Acme Analytics", title="Business Analyst (repost)",
                         discovered_at="2026-10-10T16:00:00+00:00", duplicate_of_job_id=ids["j1"],
                         job_url_canonical=None)
    _score(db, ids["j1"], 60.0, "possible_match", current=0, version="v0")  # superseded score
    _score(db, ids["j1"], 85.0, "strong_match", shortlisted=1,
           reasons=["Strong SQL match", "Python experience", "Analytics domain", "Fourth reason"],
           missing=["Tableau"])
    _score(db, ids["j2"], 72.0, "possible_match")
    _score(db, ids["j3"], 55.0, "needs_review", flags=[{
        "code": "no_sponsorship", "severity": "review", "excerpt": "We do not sponsor visas",
        "message": "Listing says the employer does not sponsor work visas."}])
    _score(db, ids["j4"], 90.0, "strong_match")

    for n in range(6, 14):
        ids[f"j{n}"] = make_job(employer=f"Employer{n}", title=f"Analyst {n}", discovered_at=old)
    ids["a_prepared"] = _app(db, ids["j6"], "A-prep", "awaiting_approval",
                             prepared_at="2026-10-10T13:00:00+00:00")
    ids["a_attempted"] = _app(db, ids["j7"], "A-att", "submission_attempted")
    ids["a_noevidence"] = _app(db, ids["j8"], "A-noev", "submitted", submitted_at="2026-10-10T12:00:00+00:00")
    ids["a_page"] = _app(db, ids["j9"], "A-page", "submitted", confirmation="page_evidence",
                         submitted_at="2026-10-10T14:00:00+00:00", notes="+1 call me")
    ids["a_user"] = _app(db, ids["j10"], "A-user", "submitted", confirmation="user_reported",
                         submitted_at="2026-10-08T14:00:00+00:00")
    ids["a_interview"] = _app(db, ids["j11"], "A-int", "interview_scheduled", confirmation="email_confirmed",
                              submitted_at="2026-10-01T14:00:00+00:00")
    _event(db, ids["a_interview"], "submitted", "2026-10-01T14:00:00+00:00", evidence={"type": "page"})
    _event(db, ids["a_interview"], "interview_requested", "2026-10-10T14:00:00+00:00", evidence={"type": "email"})
    _event(db, ids["a_interview"], "interview_scheduled", "2026-10-10T15:00:00+00:00", evidence={"type": "email"})
    ids["a_rejected"] = _app(db, ids["j12"], "A-rej", "rejected", confirmation="email_confirmed",
                             submitted_at="2026-09-30T14:00:00+00:00")
    _event(db, ids["a_rejected"], "rejected", "2026-10-05T14:00:00+00:00", evidence={"type": "email"})
    ids["a_evented"] = _app(db, ids["j13"], "A-evt", "submitted", submitted_at="2026-10-09T14:00:00+00:00")
    _event(db, ids["a_evented"], "submitted", "2026-10-09T14:00:00+00:00", evidence={"type": "other", "ref": "x"})

    db.insert("approval_requests", {
        "approval_uid": "AP-live", "request_type": "submit_application", "application_id": ids["a_prepared"],
        "summary": "Submit application: Employer6 — Analyst 6", "payload": "{}", "payload_hash": "x",
        "status": "pending", "requested_at": "2026-10-10T13:00:00+00:00", "expires_at": "2099-01-01T00:00:00+00:00"})
    db.insert("approval_requests", {
        "approval_uid": "AP-old", "request_type": "submit_application", "application_id": ids["a_attempted"],
        "summary": "Expired request", "payload": "{}", "payload_hash": "y", "status": "pending",
        "requested_at": "2020-01-01T00:00:00+00:00", "expires_at": "2020-01-04T00:00:00+00:00"})
    db.insert("followups", {
        "application_id": ids["a_user"], "followup_type": "application_no_response",
        "due_at": "2026-10-10T12:00:00+00:00", "dedupe_key": "f1", "suggested_action": "Send a polite check-in",
        "draft_text": "=cmd|' /C calc'!A0", "status": "drafted", "created_at": old, "updated_at": old})
    db.insert("followups", {
        "application_id": ids["a_page"], "followup_type": "application_no_response",
        "due_at": "2026-10-02T12:00:00+00:00", "dedupe_key": "f2", "suggested_action": "Done already",
        "status": "completed", "completed_at": old, "created_at": old, "updated_at": old})
    db.insert("errors", {
        "occurred_at": "2026-10-10T18:00:00+00:00", "component": "email", "error_kind": "authorization_required",
        "message": f"Gmail refused: Authorization: {RAW_SECRET}", "requires_human_action": 1,
        "human_action": "Run `job-agent gmail auth`."})
    email_id = db.insert("email_messages", {
        "gmail_message_id": "m1", "gmail_thread_id": "t1", "received_at": "2026-10-10T17:00:00+00:00",
        "from_address": "recruiter@globex.example", "from_domain": "globex.example",
        "subject": "<script>alert(1)</script> Quick chat?", "category": "recruiter_message",
        "processed_at": "2026-10-10T17:05:00+00:00"})
    db.insert("email_associations", {"email_id": email_id, "job_id": ids["j2"], "method": "domain",
                                     "confidence": 0.9, "created_at": "2026-10-10T17:05:00+00:00"})
    db.insert("document_versions", {
        "job_id": ids["j9"], "doc_type": "resume", "version": 1, "format": "docx",
        "file_name": "Employer9_Analyst-9_Resume_2026-10-10.docx", "relative_path": "x/resume.docx",
        "sha256": "0" * 64, "generator": "rules", "profile_hash": "p", "created_at": "2026-10-10T12:30:00+00:00"})
    db.insert("agent_runs", {
        "run_uid": "R-1", "task_name": "discover", "trigger": "scheduled", "status": "failed",
        "started_at": "2026-10-10T11:15:00+00:00", "finished_at": "2026-10-10T11:16:00+00:00",
        "error_kind": "recoverable", "error_message": "network down"})
    db.insert("audit_events", {"occurred_at": "2026-10-10T13:00:00+00:00", "actor": "system",
                               "action": "approval_requested", "entity_type": "approval", "entity_id": "AP-live",
                               "outcome": "success", "details": "{}"})
    db.execute("INSERT INTO agent_state (key, value, updated_at) VALUES ('gmail_auth_error', ?, ?)",
               ("invalid_grant: Token has been expired or revoked.", old))
    return ids


def _sheet_rows(ws, header_row=1):
    headers = [c.value for c in ws[header_row]]
    return headers, [dict(zip(headers, (c for c in row))) for row in ws.iter_rows(min_row=header_row + 1)
                     if any(c.value is not None for c in row)]


def _metric(ws, label):
    for row in ws.iter_rows(min_row=4, max_col=2):
        if row[0].value == label:
            return row[1].value
    raise AssertionError(f"metric {label!r} not found")


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
def test_all_sheets_exist_with_headers(ctx, seeded):
    result = generate_daily_report(ctx, D)
    assert result.xlsx_path.name == "Job_Search_Report_2026-10-10.xlsx"
    assert result.markdown_path.name == "Daily_Summary_2026-10-10.md"
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert wb.sheetnames == list(SHEET_NAMES) == [
        "Dashboard", "All Jobs", "Applications", "Daily Summary", "Follow-ups", "Activity Log", "Configuration"]
    expected = {
        "All Jobs": ["Employer", "Role", "Location", "Work Arrangement", "Salary Min", "Salary Max",
                     "Salary Period", "Salary Note", "Posting Date", "Discovery Date", "Fit Score",
                     "Match Category", "Missing Qualifications", "Eligibility Concerns", "Source", "Job URL",
                     "Application URL", "Status", "Last Checked (America/New_York)"],
        "Applications": ["Employer", "Job Title", "Application Date", "Date Type", "Status", "Résumé File",
                         "Cover Letter File", "Confirmation Status", "Follow-up Date", "Last Activity",
                         "Next Action", "Notes"],
        "Daily Summary": ["Section", "Item", "Details", "Link"],
        "Follow-ups": ["Employer", "Role", "Last Activity", "Follow-up Date", "Suggested Action",
                       "Draft (not sent)", "Approval Status", "Completion Status"],
        "Activity Log": ["Timestamp (America/New_York)", "Task / Actor", "Job / Application", "Attempted Action",
                         "Outcome", "Error", "Human Action Required"],
        "Configuration": ["Setting", "Value"],
    }
    for name, headers in expected.items():
        actual = [c.value for c in wb[name][1]]
        assert set(headers) <= set(actual), (name, set(headers) - set(actual))
    assert [c.value for c in wb["Dashboard"][3]][:3] == ["Metric", "Value", "How it is counted"]
    summary_sections = {r["Section"].value for r in _sheet_rows(wb["Daily Summary"])[1]}
    for section in ("New opportunities", "Top matches", "Generated documents", "Applications prepared",
                    "Actual submissions (evidence-based)", "Recruiter messages", "Interviews", "Assessments",
                    "Deadlines (overdue or within 7 days)", "Approvals pending", "Failures"):
        assert section in summary_sections


def test_counts_match_seeded_data(ctx, seeded):
    result = generate_daily_report(ctx, D)
    c = result.stats["counts"]
    assert c["total_jobs"] == 11  # j4 discovered after the day; j5 is a duplicate
    assert c["new_jobs_today"] == 2  # j1 and j2 (local Oct 10 23:30)
    assert c["jobs_scored"] == 3  # j1, j2, j3 (j4 after the day)
    assert c["shortlisted"] == 1
    assert c["applications_total"] == 8
    assert c["applications_prepared_today"] == 1
    assert c["awaiting_approval"] == 1
    assert c["approvals_pending"] == 1  # the expired one is excluded
    # page evidence, user reported, interview (email), rejected (email), event evidence.
    # NOT: awaiting approval, submission_attempted, "submitted" without any evidence.
    assert c["submitted"] == 5
    assert c["submitted_today"] == 1
    assert c["confirmations"] == 3
    assert c["interviews"] == 1 and c["rejections"] == 1 and c["offers"] == 0 and c["assessments"] == 0
    assert c["pending_followups"] == 1 and c["followups_due"] == 1
    assert c["errors_requiring_attention"] == 1
    assert c["documents_generated_today"] == 1
    assert c["recruiter_messages_today"] == 1

    ws = openpyxl.load_workbook(result.xlsx_path)["Dashboard"]
    assert _metric(ws, "Submitted (evidence-based)") == 5
    assert _metric(ws, "New jobs today") == 2
    assert _metric(ws, "Applications awaiting approval") == 1
    assert _metric(ws, "Errors requiring attention") == 1

    apps = {r["Application ID"].value: r for r in _sheet_rows(openpyxl.load_workbook(result.xlsx_path)
                                                              ["Applications"])[1]}
    assert apps["A-noev"]["Counted as Submitted"].value == "No"
    assert "not counted" in apps["A-noev"]["Submission Evidence"].value
    assert apps["A-prep"]["Date Type"].value == "Created (not submitted)"
    assert apps["A-page"]["Date Type"].value == "Submitted"
    assert apps["A-page"]["Résumé File"].value == "Employer9_Analyst-9_Resume_2026-10-10.docx"


def test_new_today_uses_local_timezone_boundaries(db, seeded):
    start, end = q.local_day_bounds(D, "America/New_York")
    assert (start, end) == ("2026-10-10T04:00:00+00:00", "2026-10-11T04:00:00+00:00")
    local = {r["title"] for r in q.new_jobs(db, start, end, 50)[0]}
    assert local == {"Business Analyst", EVIL_TITLE}
    utc_start, utc_end = q.local_day_bounds(D, "UTC")
    utc = {r["title"] for r in q.new_jobs(db, utc_start, utc_end, 50)[0]}
    assert utc == {"Business Analyst", "Data Analyst"}  # 02:00Z belongs to Oct 10 in UTC, not in New York
    # DST change (Nov 1 2026 in New York) gives a 25-hour day
    assert q.local_day_bounds(date(2026, 11, 1), "America/New_York") == (
        "2026-11-01T04:00:00+00:00", "2026-11-02T05:00:00+00:00")


def test_trend_counts(db, seeded):
    trend = q.trend_counts(db, 14, tz_name="America/New_York", end_date=D)
    assert len(trend) == 14 and trend[-1]["date"] == "2026-10-10" and trend[0]["date"] == "2026-09-27"
    by_date = {t["date"]: t for t in trend}
    assert by_date["2026-10-10"]["discovered"] == 2
    assert by_date["2026-10-09"]["discovered"] == 1  # j3 at 22:00 local
    assert by_date["2026-10-01"]["discovered"] == 8
    assert by_date["2026-10-10"]["submitted"] == 1 and by_date["2026-10-10"]["interviews"] == 1
    assert by_date["2026-10-05"]["rejections"] == 1
    assert sum(t["submitted"] for t in trend) == 5
    assert set(trend[0]) == {"date", "discovered", "applications", "submitted", "interviews", "rejections", "offers"}
    with pytest.raises(ValueError):
        q.trend_counts(db, 0)


def test_hyperlinks_only_for_http(ctx, seeded):
    result = generate_daily_report(ctx, D)
    wb = openpyxl.load_workbook(result.xlsx_path)
    _, rows = _sheet_rows(wb["All Jobs"])
    by_employer = {r["Employer"].value + "|" + str(r["Job ID"].value): r for r in rows}
    j1 = next(r for k, r in by_employer.items() if k.startswith("Acme Analytics|J-test00001"))
    assert j1["Job URL"].hyperlink.target == "https://example.com/jobs/1"
    assert j1["Job URL"].style == "Hyperlink"
    assert j1["Application URL"].hyperlink is None and j1["Application URL"].value == "javascript:alert(1)"
    j3 = next(r for r in rows if r["Employer"].value == "Initech")
    assert j3["Job URL"].hyperlink is None and j3["Job URL"].value == "ftp://initech.example/job3"
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None:
                    assert cell.hyperlink.target.startswith(("http://", "https://")), cell.hyperlink.target


def test_freeze_panes_autofilter_and_formats(ctx, seeded):
    result = generate_daily_report(ctx, D)
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert wb["Dashboard"].freeze_panes == "A4" and wb["Dashboard"].auto_filter.ref
    for name in SHEET_NAMES[1:]:
        assert wb[name].freeze_panes == "A2", name
        assert wb[name].auto_filter.ref and wb[name].auto_filter.ref.startswith("A1:"), name
    ws = wb["All Jobs"]
    _, rows = _sheet_rows(ws)
    j1 = next(r for r in rows if r["Job ID"].value == "J-test00001")
    assert j1["Salary Min"].value == 80000 and j1["Salary Min"].number_format == "$#,##0"
    assert j1["Salary Max"].number_format == "$#,##0"
    assert j1["Salary Period"].value == "per year" and "base salary" in j1["Salary Note"].value
    assert j1["Fit Score"].value == 85.0 and j1["Fit Score"].number_format == "0.0"
    assert j1["Match Category"].value == "Strong match"
    assert j1["Posting Date"].number_format == "yyyy-mm-dd"
    assert j1["Posting Date"].value.date() == date(2026, 10, 8)
    disc = j1["Discovery Date"]
    assert disc.number_format == "yyyy-mm-dd hh:mm"
    assert (disc.value.hour, disc.value.day) == (11, 10)  # 15:00Z -> 11:00 local
    assert j1["Missing Qualifications"].alignment.wrap_text
    j3 = next(r for r in rows if r["Employer"].value == "Initech")
    assert "no_sponsorship [review]" in j3["Eligibility Concerns"].value
    assert ws.conditional_formatting  # score scale + category/status fills
    rule_types = {rule.type for cf in ws.conditional_formatting for rule in cf.rules}
    assert {"colorScale", "cellIs"} <= rule_types
    for col_letter in ("A", "C"):
        assert 8 <= ws.column_dimensions[col_letter].width <= 60
    # Activity log is in local time and redacted
    _, activity = _sheet_rows(wb["Activity Log"])
    assert activity and all(r["Timestamp (America/New_York)"].number_format == "yyyy-mm-dd hh:mm"
                            for r in activity)
    assert {r["Type"].value for r in activity} == {"run", "error", "audit"}


def test_charts_on_dashboard(ctx, seeded):
    result = generate_daily_report(ctx, D)
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert len(wb["Dashboard"]._charts) == 3
    with zipfile.ZipFile(result.xlsx_path) as zf:
        charts = [n for n in zf.namelist() if n.startswith("xl/charts/chart")]
        kinds = "".join(zf.read(n).decode() for n in charts)
    assert len(charts) == 3 and "barChart" in kinds and "lineChart" in kinds


def test_formula_injection_neutralized(ctx, seeded):
    result = generate_daily_report(ctx, D)
    wb = openpyxl.load_workbook(result.xlsx_path)
    _, rows = _sheet_rows(wb["All Jobs"])
    evil = next(r for r in rows if r["Employer"].value == "Globex")["Role"]
    assert evil.data_type == "s" and evil.value == "'" + EVIL_TITLE
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@")), (ws.title, cell.coordinate)
    fups = _sheet_rows(wb["Follow-ups"])[1]
    assert fups[0]["Draft (not sent)"].value.startswith("'=cmd")
    notes = next(r for r in _sheet_rows(wb["Applications"])[1] if r["Application ID"].value == "A-page")["Notes"]
    assert notes.value == "'+1 call me"


def test_no_secrets_in_report(config, seeded, monkeypatch):
    from job_agent.config import Secrets
    from job_agent.context import AppContext
    from job_agent.security.passwords import hash_password

    pw_hash = hash_password("correct horse battery staple")
    values = {
        "JOB_AGENT_SESSION_SECRET": "sess-SECRET-9f8e7d6c5b4a",
        "JOB_AGENT_DASHBOARD_PASSWORD_HASH": pw_hash,
        "JOB_AGENT_DASHBOARD_USERNAME": "admin-user-qwerty",
        "JOB_AGENT_CONTACT_EMAIL": "private.contact@example.org",
    }
    for k, v in values.items():
        monkeypatch.setenv(k, v)
    secrets = Secrets(session_secret=values["JOB_AGENT_SESSION_SECRET"], dashboard_password_hash=pw_hash,
                      dashboard_username=values["JOB_AGENT_DASHBOARD_USERNAME"],
                      contact_email=values["JOB_AGENT_CONTACT_EMAIL"])
    c = AppContext.from_config(config, secrets)
    try:
        result = generate_daily_report(c, D)
        texts = []
        wb = openpyxl.load_workbook(result.xlsx_path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                texts += [str(cell.value) for cell in row if cell.value is not None]
        with zipfile.ZipFile(result.xlsx_path) as zf:
            texts += [zf.read(n).decode("utf-8", "replace") for n in zf.namelist()]
        texts.append(result.markdown_path.read_text(encoding="utf-8"))
        texts.append(c.db.scalar("SELECT stats FROM reports WHERE report_date = ?", (D.isoformat(),)))
        blob = "\n".join(texts)
        for secret in [*values.values(), pw_hash.split("$")[-1], RAW_SECRET.split()[1]]:
            assert secret not in blob, secret
        config_values = {str(r[0].value): r[1].value for r in wb["Configuration"].iter_rows(min_row=2)}
        assert config_values["Secrets"].startswith("not included")
        assert config_values["Operating mode"] == "supervised"
        assert config_values["App version"]
        assert config_values["scoring.shortlist_threshold"] == 70
    finally:
        c.close()


def test_regeneration_is_idempotent_and_other_dates_untouched(ctx, seeded):
    other_x, other_md = report_paths(ctx.config, date(2026, 10, 9))
    ctx.config.reports_dir.mkdir(parents=True, exist_ok=True)
    other_x.write_bytes(b"older report")
    other_md.write_text("older summary", encoding="utf-8")
    before = (other_x.stat().st_mtime_ns, other_md.stat().st_mtime_ns)

    first = generate_daily_report(ctx, D)
    second = generate_daily_report(ctx, D)
    assert (first.xlsx_path, first.markdown_path) == (second.xlsx_path, second.markdown_path)
    assert first.generation_count == 1 and second.generation_count == 2
    rows = ctx.db.query_all("SELECT * FROM reports")
    assert len(rows) == 1 and rows[0]["generation_count"] == 2
    assert rows[0]["xlsx_file"] == "Job_Search_Report_2026-10-10.xlsx"
    assert rows[0]["first_generated_at"] <= rows[0]["last_generated_at"]
    assert json.loads(rows[0]["stats"])["counts"]["submitted"] == 5
    assert other_x.read_bytes() == b"older report" and other_md.read_text(encoding="utf-8") == "older summary"
    assert (other_x.stat().st_mtime_ns, other_md.stat().st_mtime_ns) == before
    assert sorted(p.name for p in ctx.config.reports_dir.iterdir()) == [
        "Daily_Summary_2026-10-09.md", "Daily_Summary_2026-10-10.md",
        "Job_Search_Report_2026-10-09.xlsx", "Job_Search_Report_2026-10-10.xlsx"]  # no temp files left
    openpyxl.load_workbook(second.xlsx_path)  # still a valid workbook
    audits = ctx.db.scalar("SELECT COUNT(*) FROM audit_events WHERE action = 'report_generated'")
    assert audits == 2

    history = list_reports(ctx)
    assert history[0]["report_date"] == "2026-10-10" and history[0]["xlsx_file_exists"]
    assert resolve_report_file(ctx, history[0]["id"], "markdown") == second.markdown_path


def test_empty_database_produces_valid_report(ctx):
    result = generate_daily_report(ctx)  # defaults to today in app.timezone
    from job_agent.utils.time import local_today

    assert result.report_date == local_today("America/New_York")
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert wb.sheetnames == list(SHEET_NAMES)
    assert all(v == 0 for v in result.stats["counts"].values())
    assert len(wb["Dashboard"]._charts) == 3
    md = result.markdown_path.read_text(encoding="utf-8")
    assert "Nothing needs your action right now." in md and "No scored matches yet." in md
    summary = _sheet_rows(wb["Daily Summary"])[1]
    assert all(r["Item"].value == "None" for r in summary)


def test_deadlines_section(ctx, seeded):
    ctx.db.insert("followups", {
        "application_id": seeded["a_interview"], "followup_type": "assessment_due",
        "due_at": "2026-10-12T16:00:00+00:00", "dedupe_key": "f3", "suggested_action": "Finish the case study",
        "status": "pending", "created_at": "2026-10-09T00:00:00+00:00", "updated_at": "2026-10-09T00:00:00+00:00"})
    ctx.db.insert("followups", {
        "application_id": seeded["a_page"], "followup_type": "deadline", "due_at": "2026-10-30T16:00:00+00:00",
        "dedupe_key": "f4", "suggested_action": "Too far out", "status": "pending",
        "created_at": "2026-10-09T00:00:00+00:00", "updated_at": "2026-10-09T00:00:00+00:00"})
    result = generate_daily_report(ctx, D)
    rows = _sheet_rows(openpyxl.load_workbook(result.xlsx_path)["Daily Summary"])[1]
    deadlines = [(r["Item"].value, r["Details"].value) for r in rows
                 if r["Section"].value == "Deadlines (overdue or within 7 days)"]
    assert deadlines == [("Employer11 — Analyst 11", "Assessment due: 2026-10-12 12:00"),
                         ("Acme Analytics — Business Analyst", "Listing closes: 2026-10-13 19:00")]
    assert "2 deadline(s) overdue or within 7 days" in result.markdown_path.read_text(encoding="utf-8")


def test_row_caps_are_reported(ctx, seeded, monkeypatch):
    from job_agent.reporting import data as report_data

    monkeypatch.setattr(report_data, "MAX_JOB_ROWS", 2)
    result = generate_daily_report(ctx, D)
    wb = openpyxl.load_workbook(result.xlsx_path)
    ws = wb["All Jobs"]
    assert ws.auto_filter.ref == "A1:U3"
    assert [c.value for c in ws["A"] if c.value][-1].startswith("Showing 2 of 11 rows")
    assert _metric(wb["Dashboard"], "Rows not shown: All Jobs") == 9
    assert result.stats["truncated"] == {"All Jobs": {"total": 11, "shown": 2}}


def test_future_date_rejected(ctx):
    from job_agent.errors import PermanentError

    with pytest.raises(PermanentError, match="future"):
        generate_daily_report(ctx, date(2999, 1, 1))


def test_report_task_via_run_task(ctx, seeded):
    from job_agent.models import RunStatus
    from job_agent.tasks import run_task

    result = run_task(ctx, "report", trigger="scheduled")
    assert result.status == RunStatus.SUCCEEDED, result.error
    assert result.summary["xlsx_file"].startswith("Job_Search_Report_")
    assert result.summary["generation_count"] == 1
    actor = ctx.db.scalar("SELECT actor FROM audit_events WHERE action = 'report_generated'")
    assert actor == "scheduler"


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def test_markdown_summary(ctx, seeded):
    result = generate_daily_report(ctx, D)
    md = result.markdown_path.read_text(encoding="utf-8")
    assert md.startswith("# Job search summary — 2026-10-10")
    for heading in ("## Needs your action", "## Top matches", "## Documents generated", "## Applications",
                    "## Messages, interviews and assessments", "## Failures", "## Next steps"):
        assert heading in md
    # Needs your action: approvals, eligibility review, Gmail auth, errors, follow-ups due
    assert "1 approval(s) pending" in md
    assert "1 job(s) need an eligibility review" in md and "no_sponsorship \\[review\\]" in md
    assert "Gmail authorization problem" in md and "invalid_grant" in md
    assert "1 error(s) need your attention" in md
    assert "1 follow-up(s) due" in md
    # Top matches with score + up to 3 reasons (j4 was discovered after the report day)
    assert "**85.0** · Acme Analytics — Business Analyst" in md
    assert "Strong SQL match" in md and "Analytics domain" in md and "Fourth reason" not in md
    assert "[Open posting](https://example.com/jobs/1)" in md
    assert "**Actually submitted today (evidence-based): 1**" in md
    assert "Prepared today: 1 (not yet submitted: 1)" in md
    assert "- Interviews: 1\n  - Employer11 — Analyst 11 — Interview scheduled" in md  # one line per application
    assert "Recruiter messages: 1" in md
    # Untrusted text is neutralized: no raw HTML, no live Markdown link from the evil title
    assert "<script>" not in md and "&lt;script&gt;" in md
    assert '· Globex — \\=HYPERLINK("http://evil.example/x","click me")' in md
    assert RAW_SECRET.split()[1] not in md
    assert "javascript:" not in md


def test_text_helpers():
    from job_agent.reporting.text import excel_text, md_escape, md_url, safe_http_url

    assert excel_text("=1+1") == "'=1+1"
    assert excel_text("@SUM(A1)") == "'@SUM(A1)" and excel_text("-2") == "'-2" and excel_text("+x") == "'+x"
    assert excel_text("\t=x") == "'\t=x" and excel_text("Plain title") == "Plain title"
    assert excel_text("bad\x00\x07char") == "badchar"
    assert len(excel_text("x" * 40_000)) <= 32_000
    assert safe_http_url("https://a.example/x?y=1") == "https://a.example/x?y=1"
    for bad in ("javascript:alert(1)", "file:///etc/passwd", "mailto:a@b.c", "/relative", "https://a b.example",
                "http://", "https://x.example/<script>", None, 5):
        assert safe_http_url(bad) is None, bad
    assert md_escape("<b>[x](http://e.example)</b>\n# h") == "&lt;b&gt;\\[x\\]\\(http://e.example)&lt;/b&gt; # h"
    assert "](" not in md_escape("[click](https://evil.example)") and md_escape("Analyst (Remote)") == "Analyst (Remote)"
    assert md_escape("# Heading") == "\\# Heading" and md_escape("> quote") == "&gt; quote"
    assert md_escape("my_file_name.docx") == "my_file_name.docx" and md_escape("_x_ *y*") == "\\_x\\_ \\*y\\*"
    assert md_escape("`code`") == "\\`code\\`" and md_escape("![img](http://x)") == "!\\[img\\]\\(http://x)"
    assert md_escape("1. item").startswith("1\\.") and md_escape("- item").startswith("\\-")
    assert md_url("https://ok.example/a") == "[Open posting](https://ok.example/a)"
    assert md_url("javascript:x") == "" and md_url("https://x.example/a)(b") == ""


def test_counts_with_real_repo_state_machine(ctx, make_job):
    """Applications driven through database.repo (as other modules do) are counted correctly."""
    from job_agent.database import repo
    from job_agent.models import ApprovalType, PipelineState as S

    db = ctx.db

    def prepared(job_id):
        app_id = repo.create_application(db, job_id, actor="system")
        repo.transition_application(db, app_id, S.DOCUMENTS_READY, actor="system")
        repo.transition_application(db, app_id, S.AWAITING_APPROVAL, actor="system")
        return app_id

    def attempt(app_id):
        payload = {"application_id": app_id, "answers": {"q": "a"}}
        uid = repo.create_approval(db, ApprovalType.SUBMIT_APPLICATION, payload, summary="Submit", actor="system",
                                   application_id=app_id)
        repo.decide_approval(db, uid, approve=True, actor="user")
        repo.transition_application(db, app_id, S.SUBMISSION_ATTEMPTED, actor="browser", approval_uid=uid,
                                    approved_payload=payload)

    browser_ok = prepared(make_job())
    attempt(browser_ok)
    repo.transition_application(db, browser_ok, S.SUBMITTED, actor="browser",
                                evidence={"type": "page", "url": "https://example.com/thanks"})
    unknown = prepared(make_job())
    attempt(unknown)
    repo.transition_application(db, unknown, S.NEEDS_MANUAL_ACTION, actor="browser")
    waiting = prepared(make_job())
    repo.create_approval(db, ApprovalType.SUBMIT_APPLICATION, {"application_id": waiting}, summary="Submit",
                         actor="system", application_id=waiting)
    manual = repo.create_application(db, make_job(), actor="system")
    repo.transition_application(db, manual, S.DOCUMENTS_READY, actor="system")
    repo.transition_application(db, manual, S.SUBMITTED, actor="user")

    c = generate_daily_report(ctx).stats["counts"]
    assert c["submitted"] == 2 and c["submitted_today"] == 2  # browser with page evidence + user-reported
    assert c["confirmations"] == 1  # only the confirmation page; "I applied myself" is not a confirmation
    assert c["awaiting_approval"] == 1 and c["approvals_pending"] == 1
    assert c["applications_prepared_today"] == 3 and c["applications_total"] == 4
    assert c["new_jobs_today"] == 4
