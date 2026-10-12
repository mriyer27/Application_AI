"""Builds the daily Excel workbook (OpenPyXL) from a :class:`ReportData` snapshot.

Security: every string that comes from the database goes through
:func:`text.excel_text` (formula-injection neutralization, illegal-character
removal); only plain http(s) URLs become hyperlinks; error/audit text is
redacted. No secrets are ever read here.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Iterable, NamedTuple, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from job_agent.database.connection import from_json
from job_agent.reporting import queries as q
from job_agent.reporting import text as t
from job_agent.reporting.data import ReportData
from job_agent.utils.time import utcnow

SHEET_NAMES: tuple[str, ...] = (
    "Dashboard", "All Jobs", "Applications", "Daily Summary", "Follow-ups", "Activity Log", "Configuration",
)

DATETIME_FMT = "yyyy-mm-dd hh:mm"
DATE_FMT = "yyyy-mm-dd"
SALARY_FMT = "$#,##0"
SCORE_FMT = "0.0"

_HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=12)
_TOP = Alignment(vertical="top")
_WRAP = Alignment(vertical="top", wrap_text=True)

_GREEN, _BLUE, _YELLOW, _ORANGE, _RED, _GREY = "C6EFCE", "DDEBF7", "FFEB9C", "F8CBAD", "FFC7CE", "D9D9D9"

#: Conditional-format fills for status labels (as shown in the sheets).
STATUS_FILLS: dict[str, str] = {
    **{t.status_label(s): _GREEN for s in ("submitted", "confirmation_received", "offer_received")},
    **{t.status_label(s): _BLUE for s in ("interview_requested", "interview_scheduled", "assessment_received")},
    **{t.status_label(s): _YELLOW for s in ("shortlisted", "documents_pending", "documents_ready", "in_progress",
                                             "awaiting_approval", "pending", "drafted", "approved")},
    **{t.status_label(s): _ORANGE for s in ("submission_attempted", "needs_manual_action")},
    **{t.status_label(s): _RED for s in ("failed",)},
    **{t.status_label(s): _GREY for s in ("rejected", "withdrawn", "closed", "filtered_out", "dismissed",
                                           "completed")},
}
CATEGORY_FILLS: dict[str, str] = {
    "Strong match": _GREEN, "Possible match": _YELLOW, "Weak match": _GREY, "Needs review": _ORANGE,
}


class Link(NamedTuple):
    url: Any
    text: str | None = None


class Money(NamedTuple):
    amount: float | None
    currency: str | None


@dataclass(frozen=True)
class Col:
    header: str
    kind: str = "text"  # text | wrap | datetime | date | score | money | int | link
    width: int = 18  # maximum column width (characters)


# --------------------------------------------------------------------------- #
# Low-level cell writing
# --------------------------------------------------------------------------- #
def _money_format(currency: str | None) -> str:
    code = (currency or "USD").strip().upper()
    if code in {"USD", "$"}:
        return SALARY_FMT
    code = "".join(ch for ch in code if ch.isalpha())[:3] or "?"
    return f'#,##0 "{code}"'


def _put(ws: Worksheet, row: int, column: int, value: Any, col: Col) -> None:
    cell = ws.cell(row=row, column=column)
    if value is None or value == "":
        return
    if isinstance(value, Link):
        url = t.safe_http_url(value.url)
        if url is not None:
            cell.value = t.excel_text(value.text or url)
            cell.hyperlink = url
            cell.style = "Hyperlink"
        else:  # not a plain http(s) URL: plain text only, never clickable
            cell.value = t.excel_text(str(value.text or value.url))
        cell.alignment = _TOP
        return
    if isinstance(value, Money):
        if value.amount is None or not math.isfinite(float(value.amount)):
            return
        cell.value = float(value.amount)
        cell.number_format = _money_format(value.currency)
    elif isinstance(value, datetime):
        cell.value = value.replace(tzinfo=None)
        cell.number_format = DATETIME_FMT
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = DATE_FMT
    elif isinstance(value, bool):
        cell.value = value
    elif isinstance(value, (int, float)):
        cell.value = value
        if col.kind == "score":
            cell.number_format = SCORE_FMT
    else:
        cell.value = t.excel_text(str(value))
    cell.alignment = _WRAP if col.kind == "wrap" else _TOP


def _display_len(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime):
        return 16
    if isinstance(value, date):
        return 10
    if isinstance(value, Link):
        return len(str(value.text or value.url or ""))
    if isinstance(value, Money):
        return 10
    if isinstance(value, float):
        return 8
    return max((len(line) for line in str(value).splitlines()), default=0)


def _write_table(ws: Worksheet, cols: Sequence[Col], rows: Iterable[Sequence[Any]], *, header_row: int = 1,
                 freeze: bool = True, note: str | None = None) -> int:
    """Header + rows with formats, freeze panes, auto-filter and capped widths. Returns the last data row."""
    for ci, col in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=ci, value=col.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [len(c.header) for c in cols]
    r = header_row
    for r_offset, values in enumerate(rows, start=1):
        r = header_row + r_offset
        for ci, (col, value) in enumerate(zip(cols, values), start=1):
            _put(ws, r, ci, value, col)
            if r_offset <= 500:
                widths[ci - 1] = max(widths[ci - 1], _display_len(value))
    last_col = get_column_letter(len(cols))
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(r, header_row + 1)}"
    if freeze:
        ws.freeze_panes = f"A{header_row + 1}"
    for ci, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = max(8, min(col.width, widths[ci - 1] + 2))
    if note:
        cell = ws.cell(row=r + 2, column=1, value=t.excel_text(note))
        cell.font = Font(italic=True)
    return r


def _status_rules(ws: Worksheet, column: int, first_row: int, last_row: int) -> None:
    letter = get_column_letter(column)
    rng = f"{letter}{first_row}:{letter}{max(first_row, last_row)}"
    for label, color in STATUS_FILLS.items():
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))


def _category_rules(ws: Worksheet, column: int, first_row: int, last_row: int) -> None:
    letter = get_column_letter(column)
    rng = f"{letter}{first_row}:{letter}{max(first_row, last_row)}"
    for label, color in CATEGORY_FILLS.items():
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill))


def _score_scale(ws: Worksheet, column: int, first_row: int, last_row: int) -> None:
    letter = get_column_letter(column)
    ws.conditional_formatting.add(
        f"{letter}{first_row}:{letter}{max(first_row, last_row)}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                       mid_type="num", mid_value=60, mid_color="FFEB84",
                       end_type="num", end_value=100, end_color="63BE7B"),
    )


def _truncation_note(data: ReportData, sheet: str) -> str | None:
    info = data.truncated.get(sheet)
    if info is None:
        return None
    return f"Showing {info['shown']} of {info['total']} rows (row cap). Use the dashboard or CLI for the rest."


# --------------------------------------------------------------------------- #
# Shared labels
# --------------------------------------------------------------------------- #
def job_label(row: dict[str, Any]) -> str:
    employer = t.clean_inline(row.get("employer") or "Unknown employer", max_len=80)
    title = t.clean_inline(row.get("title") or "Unknown role", max_len=120)
    return f"{employer} — {title}"


def _score_text(row: dict[str, Any]) -> str:
    score = row.get("overall_score")
    if score is None:
        return "not scored"
    return f"score {float(score):.1f} ({t.category_label(row.get('category'))})"


def reasons(row: dict[str, Any], n: int = 3) -> list[str]:
    items = t.text_list(row.get("reasons_to_apply")) or t.text_list(row.get("matching_qualifications"))
    return [t.clean_inline(x, max_len=160) for x in items[:n]]


def submission_evidence(row: dict[str, Any]) -> str:
    if row.get("evidenced_submitted"):
        status = row.get("confirmation_status") or "none"
        return t.evidence_label(status) if status != "none" else "evidence recorded on status change"
    if row.get("status") in q.POST_SUBMIT_STATES or row.get("submitted_at"):
        return "no evidence recorded — not counted as submitted"
    return "not submitted"


# --------------------------------------------------------------------------- #
# Sheets
# --------------------------------------------------------------------------- #
METRICS: tuple[tuple[str, str, str], ...] = (
    ("Total jobs discovered", "total_jobs",
     "Distinct (non-duplicate) listings discovered up to the end of the report day."),
    ("New jobs today", "new_jobs_today",
     "Listings first discovered during the report day (local time), including ones your filters removed."),
    ("New jobs today passing filters", "new_relevant_today",
     "New listings today that passed your search filters (not filtered out)."),
    ("Jobs scored", "jobs_scored", "Listings with a current fit score."),
    ("Shortlisted", "shortlisted", "Listings currently in the 'shortlisted' state."),
    ("Applications (total)", "applications_total", "Application records created up to the end of the report day."),
    ("Applications prepared today", "applications_prepared_today",
     "Forms prepared or approval requested today. Preparing is NOT submitting."),
    ("Applications awaiting approval", "awaiting_approval", "Applications currently waiting for your approval."),
    ("Approval requests pending", "approvals_pending", "Pending, unexpired approval requests of any type (now)."),
    ("Submitted (evidence-based)", "submitted",
     "Status 'submitted' or a later outcome AND a confirmation page/email or your own report. Prepared forms, "
     "pending approvals and attempts with unknown outcome never count."),
    ("Submitted today", "submitted_today", "Evidence-based submissions whose submission time is in the report day."),
    ("Confirmations", "confirmations",
     "Submitted applications confirmed by a confirmation page or an employer email."),
    ("Interviews", "interviews", "Applications that reached 'interview requested' or 'interview scheduled'."),
    ("Assessments", "assessments", "Applications that reached 'assessment received'."),
    ("Rejections", "rejections", "Applications that reached 'rejected'."),
    ("Offers", "offers", "Applications that reached 'offer received'."),
    ("Pending follow-ups", "pending_followups", "Follow-ups not completed or dismissed (now)."),
    ("Follow-ups due", "followups_due", "Open follow-ups due by the end of the report day."),
    ("Errors requiring attention", "errors_requiring_attention",
     "Unresolved errors that need a human action (now)."),
    ("Errors today", "errors_today", "Errors recorded during the report day."),
    ("Documents generated today", "documents_generated_today", "Résumé / cover-letter files created today."),
    ("Emails processed today", "emails_today", "Gmail messages received (or processed) during the report day."),
)


def _bar(title: str, x_title: str, y_title: str) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.legend = None
    chart.height, chart.width = 7.5, 16
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    return chart


def _dashboard(ws: Worksheet, data: ReportData) -> None:
    ws["A1"] = f"Job Search Report — {data.report_date.isoformat()}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = (f"Generated {data.generated_at.strftime('%Y-%m-%d %H:%M')} ({data.tz_name}) · job-agent "
                f"{data.app_version} · all times local · definitions in docs/reports.md")
    ws["A2"].font = Font(italic=True)
    cols = (Col("Metric", width=34), Col("Value", "int", width=14), Col("How it is counted", "wrap", width=80))
    rows: list[Sequence[Any]] = [("Report date", data.report_date, f"Local calendar day in {data.tz_name}.")]
    rows += [(label, data.counts.get(key, 0), definition) for label, key, definition in METRICS]
    for sheet, info in sorted(data.truncated.items()):
        rows.append((f"Rows not shown: {sheet}", info["total"] - info["shown"],
                     f"Only the first {info['shown']} of {info['total']} rows are included (row cap)."))
    last = _write_table(ws, cols, rows, header_row=3)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 80

    # Chart data blocks (below the metrics, columns A-B) -------------------
    row = last + 3
    ws.cell(row=row, column=1, value="Applications by status").font = _SECTION_FONT
    status_rows = [(t.status_label(s), n) for s, n in data.apps_by_status] or [("No applications yet", 0)]
    ws.cell(row=row + 1, column=1, value="Status").font = Font(bold=True)
    ws.cell(row=row + 1, column=2, value="Applications").font = Font(bold=True)
    for i, (label, n) in enumerate(status_rows, start=row + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=n)
    status_first, status_last = row + 2, row + 1 + len(status_rows)
    chart = _bar("Applications by status", "Status", "Applications")
    chart.add_data(Reference(ws, min_col=2, min_row=row + 1, max_row=status_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=status_first, max_row=status_last))
    ws.add_chart(chart, "E3")

    row = status_last + 3
    ws.cell(row=row, column=1, value="Score distribution (current scores)").font = _SECTION_FONT
    ws.cell(row=row + 1, column=1, value="Score range").font = Font(bold=True)
    ws.cell(row=row + 1, column=2, value="Jobs").font = Font(bold=True)
    for i, (label, n) in enumerate(data.score_buckets, start=row + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=n)
    score_last = row + 1 + len(data.score_buckets)
    chart = _bar("Score distribution", "Fit score", "Jobs")
    chart.add_data(Reference(ws, min_col=2, min_row=row + 1, max_row=score_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=row + 2, max_row=score_last))
    ws.add_chart(chart, "E19")

    row = max(score_last + 3, 52)  # keep the wide trend table clear of the floating charts
    ws.cell(row=row, column=1, value=f"Last {len(data.trend)} days").font = _SECTION_FONT
    headers = ("Date", "Discovered", "Applications", "Submitted", "Interviews", "Rejections", "Offers")
    for ci, h in enumerate(headers, start=1):
        ws.cell(row=row + 1, column=ci, value=h).font = Font(bold=True)
    for i, day in enumerate(data.trend, start=row + 2):
        cell = ws.cell(row=i, column=1, value=date.fromisoformat(day["date"]))
        cell.number_format = DATE_FMT
        for ci, key in enumerate(headers[1:], start=2):
            ws.cell(row=i, column=ci, value=int(day[key.lower()]))
    trend_last = row + 1 + len(data.trend)
    line = LineChart()
    line.title = f"Jobs discovered per day (last {len(data.trend)} days)"
    line.y_axis.title = "Jobs"
    line.x_axis.title = "Date"
    line.x_axis.number_format = "mm-dd"
    line.legend = None
    line.height, line.width = 7.5, 16
    line.x_axis.delete = False
    line.y_axis.delete = False
    line.add_data(Reference(ws, min_col=2, min_row=row + 1, max_row=trend_last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=row + 2, max_row=trend_last))
    ws.add_chart(line, "E35")


def _salary_note(row: dict[str, Any]) -> str:
    parts = []
    basis = row.get("salary_is_base")
    if row.get("salary_min") is not None or row.get("salary_max") is not None:
        parts.append("base salary" if basis == 1 else "total compensation" if basis == 0 else "basis not stated")
        currency = (row.get("salary_currency") or "").strip().upper()
        if not currency:
            parts.append("currency not stated")
        elif currency != "USD":
            parts.append(f"currency {t.clean_inline(currency, max_len=8)}")
    elif row.get("salary_text"):
        parts.append(f"as listed: {t.clean_inline(row['salary_text'], max_len=120)}")
    return "; ".join(parts)


def _period(value: Any) -> str:
    return "" if not value or value == "unknown" else f"per {value}"


def _all_jobs(ws: Worksheet, data: ReportData) -> None:
    tz = data.tz_name
    cols = (
        Col("Job ID", width=14), Col("Employer", width=28), Col("Role", "wrap", 40), Col("Location", "wrap", 28),
        Col("Work Arrangement", width=12), Col("Salary Min", "money", 12), Col("Salary Max", "money", 12),
        Col("Salary Period", width=12), Col("Salary Note", "wrap", 30), Col("Posting Date", "date", 11),
        Col("Discovery Date", "datetime", 16), Col("Fit Score", "score", 9), Col("Match Category", width=15),
        Col("Missing Qualifications", "wrap", 50), Col("Eligibility Concerns", "wrap", 50),
        Col("Source", width=20), Col("Job URL", "link", 40), Col("Application URL", "link", 40),
        Col("Status", width=20), Col("Listing", width=9), Col(f"Last Checked ({tz})", "datetime", 16),
    )
    rows = []
    for j in data.jobs:
        score = j.get("overall_score")
        rows.append((
            j.get("job_uid"), j.get("employer"), j.get("title"), j.get("location"),
            t.humanize(j.get("work_arrangement")),
            Money(j.get("salary_min"), j.get("salary_currency")), Money(j.get("salary_max"), j.get("salary_currency")),
            _period(j.get("salary_period")), _salary_note(j),
            t.local_date(j.get("posted_at"), tz), t.local_naive(j.get("discovered_at"), tz),
            None if score is None else round(float(score), 1), t.category_label(j.get("category")),
            t.join_list(j.get("missing_qualifications") or [], sep="\n"), t.flags_text(j.get("eligibility_flags")),
            f"{j.get('source_name')}:{j.get('source_identifier')}",
            Link(j.get("job_url")), Link(j.get("application_url")),
            t.status_label(j.get("application_status") or j.get("pipeline_state")),
            t.humanize(j.get("listing_status")), t.local_naive(j.get("last_checked_at"), tz),
        ))
    last = _write_table(ws, cols, rows, note=_truncation_note(data, "All Jobs"))
    _score_scale(ws, 12, 2, last)
    _category_rules(ws, 13, 2, last)
    _status_rules(ws, 19, 2, last)


def _applications(ws: Worksheet, data: ReportData) -> None:
    tz = data.tz_name
    cols = (
        Col("Application ID", width=14), Col("Employer", width=28), Col("Job Title", "wrap", 40),
        Col("Application Date", "datetime", 16), Col("Date Type", width=26), Col("Status", width=22),
        Col("Counted as Submitted", width=10), Col("Submission Evidence", "wrap", 30),
        Col("Résumé File", "wrap", 40), Col("Cover Letter File", "wrap", 40),
        Col("Confirmation Status", width=20), Col("Follow-up Date", "datetime", 16),
        Col("Last Activity", "datetime", 16), Col("Next Action", "wrap", 40), Col("Notes", "wrap", 50),
        Col("Job URL", "link", 40),
    )
    rows = []
    for a in data.applications:
        if a.get("submitted_at"):
            when, kind = a["submitted_at"], "Submitted"
        else:
            when, kind = a.get("created_at"), "Created (not submitted)"
        rows.append((
            a.get("application_uid"), a.get("employer"), a.get("title"), t.local_naive(when, tz), kind,
            t.status_label(a.get("status")), "Yes" if a.get("evidenced_submitted") else "No",
            submission_evidence(a), a.get("resume_file"), a.get("cover_letter_file"),
            t.evidence_label(a.get("confirmation_status")),
            t.local_naive(a.get("next_followup_due") or a.get("follow_up_due_at"), tz),
            t.local_naive(a.get("last_activity_at"), tz), a.get("next_action"),
            t.clean_inline(a.get("notes"), max_len=2000) if a.get("notes") else None, Link(a.get("job_url")),
        ))
    last = _write_table(ws, cols, rows, note=_truncation_note(data, "Applications"))
    _status_rules(ws, 6, 2, last)


def summary_rows(data: ReportData) -> list[tuple[str, str, str, Any]]:
    """(section, item, details, url) rows for the Daily Summary sheet."""
    tz = data.tz_name
    out: list[tuple[str, str, str, Any]] = []

    def section(name: str, items: list[tuple[str, str, Any]], more: int = 0) -> None:
        if not items:
            out.append((name, "None", "", None))
        for item, details, url in items:
            out.append((name, item, details, url))
        if more > 0:
            out.append((name, f"…and {more} more", "See the All Jobs / Applications sheets.", None))

    section("New opportunities", [
        (job_label(j), "; ".join(x for x in (_score_text(j), t.clean_inline(j.get("location"), max_len=60),
                                             t.humanize(j.get("work_arrangement"))) if x), j.get("job_url"))
        for j in data.new_jobs[:50]
    ], more=data.new_jobs_total - min(len(data.new_jobs), 50))
    section("Top matches", [
        (f"#{i} {job_label(j)}",
         f"{_score_text(j)}. Reasons: {'; '.join(reasons(j)) or 'none recorded'}. "
         f"Status: {t.status_label(j.get('application_status') or j.get('pipeline_state'))}.", j.get("job_url"))
        for i, j in enumerate(data.top_matches, start=1)
    ])
    section("Generated documents", [
        (t.clean_inline(d.get("file_name"), max_len=120),
         f"{t.humanize(d.get('doc_type'))} v{d.get('version')} ({d.get('format')}) for {job_label(d)}; "
         f"status {t.humanize(d.get('status'))}; generator {t.clean_inline(d.get('generator'), max_len=40)}", None)
        for d in data.documents
    ])
    section("Applications prepared", [
        (job_label(a), f"Status: {t.status_label(a.get('status'))}; prepared "
                       f"{t.local_text(a.get('prepared_at') or a.get('status_updated_at'), tz)}; "
                       f"{'submitted' if a.get('evidenced_submitted') else 'NOT submitted'}", a.get("job_url"))
        for a in data.prepared
    ])
    section("Actual submissions (evidence-based)", [
        (job_label(a), f"Submitted {t.local_text(a.get('submitted_at') or a.get('status_updated_at'), tz)}; "
                       f"evidence: {submission_evidence(a)}", a.get("job_url"))
        for a in data.submissions
    ])
    section("Confirmations", [
        (job_label(a), f"Confirmation: {t.evidence_label(a.get('confirmation_status'))}", a.get("job_url"))
        for a in data.confirmations
    ])

    def email_item(e: dict[str, Any]) -> tuple[str, str, Any]:
        bits = [f"from {t.clean_inline(e.get('from_domain') or 'unknown domain', max_len=80)}",
                t.local_text(e.get("at"), tz)]
        if e.get("employer") or e.get("title"):
            bits.append(f"re: {job_label(e)}")
        if e.get("summary"):
            bits.append(t.redacted(e.get("summary"), max_len=300))
        if e.get("injection_suspected"):
            bits.append("WARNING: possible prompt-injection content; treat as untrusted")
        return t.redacted(e.get("subject") or "(no subject)", max_len=200), "; ".join(b for b in bits if b), None

    def milestone_item(m: dict[str, Any]) -> tuple[str, str, Any]:
        details = f"{t.status_label(m.get('milestone'))} on {t.local_text(m.get('at'), tz)}"
        return job_label(m), details, m.get("job_url")

    section("Recruiter messages", [email_item(e) for e in data.emails_in(q.RECRUITER_EMAIL_CATEGORIES)])
    section("Interviews", [milestone_item(m) for m in data.interview_milestones]
            + [email_item(e) for e in data.emails_in(q.INTERVIEW_EMAIL_CATEGORIES)])
    section("Assessments", [milestone_item(m) for m in data.assessment_milestones]
            + [email_item(e) for e in data.emails_in(q.ASSESSMENT_EMAIL_CATEGORIES)])
    section("Rejections and offers", [milestone_item(m) for m in data.offer_milestones + data.rejection_milestones])
    section("Deadlines (overdue or within 7 days)", [
        (job_label(d) if (d.get("employer") or d.get("title")) else "Follow-up",
         f"{t.humanize(d.get('what'))}: {t.local_text(d.get('at'), tz)}", d.get("job_url"))
        for d in data.deadlines
    ])
    section("Approvals pending", [
        (t.clean_inline(a.get("summary"), max_len=200),
         f"{t.humanize(a.get('request_type'))} ({a.get('approval_uid')}); requested "
         f"{t.local_text(a.get('requested_at'), tz)}; expires {t.local_text(a.get('expires_at'), tz) or 'never'}"
         + (f"; {job_label(a)}" if a.get("employer") else ""), None)
        for a in data.approvals
    ])
    section("Failures", [
        (t.clean_inline(f.get("component"), max_len=80),
         f"{t.humanize(f.get('kind'))} at {t.local_text(f.get('at'), tz)}: {t.redacted(f.get('message'), max_len=400)}"
         + (f" -> {t.redacted(f.get('human_action'), max_len=200)}" if f.get("human_action") else ""), None)
        for f in data.failures
    ])
    return out


def _daily_summary(ws: Worksheet, data: ReportData) -> None:
    cols = (Col("Section", width=28), Col("Item", "wrap", 50), Col("Details", "wrap", 80), Col("Link", "link", 40))
    rows = [(s, item, details, Link(url) if url else None) for s, item, details, url in summary_rows(data)]
    _write_table(ws, cols, rows)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)


def _approval_status(f: dict[str, Any]) -> str:
    status = f.get("approval_status")
    if not status:
        return "Not requested"
    if status == "pending":
        exp = t.parse_ts(f.get("approval_expires_at"))
        if exp is not None and exp <= utcnow():
            return "Expired"
    return t.humanize(status)


def _followups(ws: Worksheet, data: ReportData) -> None:
    tz = data.tz_name
    cols = (
        Col("Employer", width=28), Col("Role", "wrap", 40), Col("Type", width=22),
        Col("Last Activity", "datetime", 16), Col("Follow-up Date", "datetime", 16),
        Col("Suggested Action", "wrap", 50), Col("Draft (not sent)", "wrap", 60), Col("Approval Status", width=16),
        Col("Completion Status", width=16), Col("Completed At", "datetime", 16),
    )
    rows = [(
        f.get("employer"), f.get("title"), t.humanize(f.get("followup_type")),
        t.local_naive(f.get("last_activity_at"), tz), t.local_naive(f.get("due_at"), tz),
        f.get("suggested_action"), (f.get("draft_text") or "")[:3000] or None, _approval_status(f),
        t.status_label(f.get("status")), t.local_naive(f.get("completed_at"), tz),
    ) for f in data.followups]
    last = _write_table(ws, cols, rows, note=_truncation_note(data, "Follow-ups"))
    _status_rules(ws, 9, 2, last)


def _audit_details(raw: Any) -> str:
    details = from_json(raw, None) if isinstance(raw, str) else raw
    if not details:
        return ""
    if isinstance(details, dict):
        text = "; ".join(f"{k}={v}" for k, v in sorted(details.items()) if v not in (None, "", [], {}))
    else:
        text = str(details)
    return t.redacted(text, max_len=500)


def _activity(ws: Worksheet, data: ReportData) -> None:
    tz = data.tz_name
    cols = (
        Col(f"Timestamp ({tz})", "datetime", 16), Col("Type", width=8), Col("Task / Actor", width=22),
        Col("Job / Application", width=20), Col("Attempted Action", "wrap", 36), Col("Outcome", width=20),
        Col("Error", "wrap", 60), Col("Human Action Required", "wrap", 40), Col("Details", "wrap", 50),
    )
    rows = [(
        t.local_naive(a.get("at"), tz), a.get("source"), a.get("task"), a.get("entity"), a.get("action"),
        t.humanize(a.get("outcome")), t.redacted(a.get("error"), max_len=1000) if a.get("error") else None,
        t.redacted(a.get("human_action"), max_len=300) if a.get("human_action") else None,
        _audit_details(a.get("details")) or None,
    ) for a in data.activity]
    _write_table(ws, cols, rows, note=_truncation_note(data, "Activity Log"))


def _configuration(ws: Worksheet, data: ReportData) -> None:
    cols = (Col("Setting", width=50), Col("Value", "wrap", 80))
    rows: list[tuple[str, Any]] = [
        ("Operating mode", data.mode),
        ("App version", data.app_version),
        ("Secrets", "not included (secrets live only in the environment / .env and are never read by reports)"),
        ("Config hash", data.config_hash),
        ("Using example config", "yes" if data.using_example_config else "no"),
        ("Report time zone", data.tz_name),
    ]
    for key, value in data.config_rows:
        rows.append((key, t.redacted(value, max_len=2000, pii=False) if isinstance(value, str) else value))
    _write_table(ws, cols, rows)


def build_workbook(data: ReportData) -> Workbook:
    wb = Workbook()
    builders = (_dashboard, _all_jobs, _applications, _daily_summary, _followups, _activity, _configuration)
    first = True
    for name, builder in zip(SHEET_NAMES, builders):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        builder(ws, data)
    wb.properties.title = f"Job Search Report {data.report_date.isoformat()}"
    wb.properties.creator = "job-agent"
    return wb
