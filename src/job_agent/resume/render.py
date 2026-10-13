"""Rendering: ATS-friendly DOCX, PDF (base-14 Helvetica) and a Markdown preview.

Layout rules: single column, standard headings (Education / Experience /
Projects / Skills / Certifications), no tables, images, text boxes, headers or
footers. Every renderer returns bytes/str; writing files (never overwriting) is
done by ``resume.repository``.
"""

from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from job_agent.errors import PermanentError
from job_agent.resume.models import CoverLetterContent, EducationEntry, ExperienceEntry, ResumeContent
from job_agent.resume.text import to_latin1

FONT_DOCX = "Calibri"
BODY_PT = 10.5


# --------------------------------------------------------------------------- #
# Shared line building (identical wording across formats)
# --------------------------------------------------------------------------- #
def _join(parts: list[str | None], sep: str = " | ") -> str:
    return sep.join(p for p in parts if p)


def education_lines(e: EducationEntry) -> tuple[str, str, str]:
    """(bold degree line, institution line, dates/GPA line)."""
    return (e.degree_line, _join([e.institution, e.location], ", "),
            _join([e.dates, f"GPA: {e.gpa}" if e.gpa else None]))


def entry_lines(x: ExperienceEntry) -> tuple[str, str]:
    """(bold title line, detail line)."""
    if x.kind == "project":
        return x.title, x.dates or ""
    return _join([x.title, x.organization], ", "), _join([x.team, x.location, x.dates])


# --------------------------------------------------------------------------- #
# DOCX (python-docx)
# --------------------------------------------------------------------------- #
def _docx_base(title: str, author: str):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_DOCX
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(2)
    for name, size in (("Heading 1", 12), ("Title", 16)):
        style = doc.styles[name]
        style.font.name = FONT_DOCX
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:  # theme fonts would override the explicit font in Word
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(attr), None)
            rfonts.set(qn("w:ascii"), FONT_DOCX)
            rfonts.set(qn("w:hAnsi"), FONT_DOCX)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(2)
    props = doc.core_properties
    props.title = title
    props.author = author
    props.comments = ""
    props.keywords = ""
    props.last_modified_by = author
    return doc


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_name_block(doc, content_header) -> None:
    from docx.shared import Pt

    name = doc.add_paragraph()
    run = name.add_run(content_header.full_name)
    run.bold = True
    run.font.size = Pt(16)
    contact = " | ".join(content_header.contact_parts())
    if contact:
        doc.add_paragraph(contact)


def resume_docx(content: ResumeContent) -> bytes:
    doc = _docx_base(f"Resume - {content.header.full_name}", content.header.full_name)
    _docx_name_block(doc, content.header)
    if content.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(content.summary)
    if content.education:
        doc.add_heading("Education", level=1)
        for e in content.education:
            degree, inst, dates = education_lines(e)
            p = doc.add_paragraph()
            p.add_run(degree).bold = True
            doc.add_paragraph(inst)
            if dates:
                doc.add_paragraph(dates)
            for b in e.bullets:
                doc.add_paragraph(b.text, style="List Bullet")
    for heading, entries in (("Experience", content.experience), ("Projects", content.projects)):
        if not entries:
            continue
        doc.add_heading(heading, level=1)
        for x in entries:
            title, detail = entry_lines(x)
            p = doc.add_paragraph()
            p.add_run(title).bold = True
            if detail:
                doc.add_paragraph(detail)
            for b in x.bullets:
                doc.add_paragraph(b.text, style="List Bullet")
    if content.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(content.skills))
    if content.certifications:
        doc.add_heading("Certifications", level=1)
        for c in content.certifications:
            doc.add_paragraph(c, style="List Bullet")
    return _docx_bytes(doc)


def cover_letter_docx(content: CoverLetterContent) -> bytes:
    doc = _docx_base(f"Cover Letter - {content.header.full_name}", content.header.full_name)
    _docx_name_block(doc, content.header)
    doc.add_paragraph(content.date_text)
    doc.add_paragraph(content.greeting)
    for para in content.paragraphs:
        doc.add_paragraph(para)
    doc.add_paragraph(content.closing)
    doc.add_paragraph(content.signature)
    return _docx_bytes(doc)


# --------------------------------------------------------------------------- #
# PDF (reportlab platypus, base-14 Helvetica)
# --------------------------------------------------------------------------- #
def _pdf_text(text: str) -> str:
    """Latin-1 transliteration + XML escaping for reportlab Paragraph markup."""
    return escape(to_latin1(text))


def _pdf_styles():
    from reportlab.lib.styles import ParagraphStyle

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=BODY_PT, leading=BODY_PT * 1.3, spaceAfter=2)
    return {
        "name": ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=16, leading=20,
                               spaceAfter=2),
        "heading": ParagraphStyle("heading", parent=body, fontName="Helvetica-Bold", fontSize=12, leading=15,
                                  spaceBefore=8, spaceAfter=3),
        "bold": ParagraphStyle("bold", parent=body, fontName="Helvetica-Bold"),
        "body": body,
        "bullet": ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4),
        "para": ParagraphStyle("para", parent=body, spaceAfter=8, leading=BODY_PT * 1.4),
    }


def _pdf_build(story: list, title: str, author: str) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch, title=to_latin1(title),
                            author=to_latin1(author), subject="", creator="job-agent", keywords="")
    doc.build(story)
    return buf.getvalue()


def resume_pdf(content: ResumeContent) -> bytes:
    from reportlab.platypus import Paragraph

    st = _pdf_styles()
    story: list = [Paragraph(_pdf_text(content.header.full_name), st["name"])]
    contact = " | ".join(content.header.contact_parts())
    if contact:
        story.append(Paragraph(_pdf_text(contact), st["body"]))

    def bullets(items) -> None:
        for b in items:
            story.append(Paragraph(_pdf_text(b.text), st["bullet"], bulletText="-"))

    if content.summary:
        story += [Paragraph("Summary", st["heading"]), Paragraph(_pdf_text(content.summary), st["body"])]
    if content.education:
        story.append(Paragraph("Education", st["heading"]))
        for e in content.education:
            degree, inst, dates = education_lines(e)
            story += [Paragraph(_pdf_text(degree), st["bold"]), Paragraph(_pdf_text(inst), st["body"])]
            if dates:
                story.append(Paragraph(_pdf_text(dates), st["body"]))
            bullets(e.bullets)
    for heading, entries in (("Experience", content.experience), ("Projects", content.projects)):
        if not entries:
            continue
        story.append(Paragraph(heading, st["heading"]))
        for x in entries:
            title, detail = entry_lines(x)
            story.append(Paragraph(_pdf_text(title), st["bold"]))
            if detail:
                story.append(Paragraph(_pdf_text(detail), st["body"]))
            bullets(x.bullets)
    if content.skills:
        story += [Paragraph("Skills", st["heading"]), Paragraph(_pdf_text(", ".join(content.skills)), st["body"])]
    if content.certifications:
        story.append(Paragraph("Certifications", st["heading"]))
        for c in content.certifications:
            story.append(Paragraph(_pdf_text(c), st["bullet"], bulletText="-"))
    return _pdf_build(story, f"Resume - {content.header.full_name}", content.header.full_name)


def cover_letter_pdf(content: CoverLetterContent) -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    st = _pdf_styles()
    story: list = [Paragraph(_pdf_text(content.header.full_name), st["name"])]
    contact = " | ".join(content.header.contact_parts())
    if contact:
        story.append(Paragraph(_pdf_text(contact), st["body"]))
    story += [Spacer(1, 12), Paragraph(_pdf_text(content.date_text), st["para"]),
              Paragraph(_pdf_text(content.greeting), st["para"])]
    story += [Paragraph(_pdf_text(p), st["para"]) for p in content.paragraphs]
    story += [Paragraph(_pdf_text(content.closing), st["body"]), Paragraph(_pdf_text(content.signature), st["body"])]
    return _pdf_build(story, f"Cover Letter - {content.header.full_name}", content.header.full_name)


# --------------------------------------------------------------------------- #
# Markdown preview (phone-friendly)
# --------------------------------------------------------------------------- #
_MD_SPECIAL = re.compile(r"([\\`*_{}\[\]<>#!])")  # no tables are used, so "|" is safe


def md_escape(text: str) -> str:
    out = _MD_SPECIAL.sub(r"\\\1", text)
    return re.sub(r"^(\s*)([-+]|\d+\.)(\s)", r"\1\\\2\3", out)


def resume_md(content: ResumeContent) -> str:
    lines = [f"# {md_escape(content.header.full_name)}", ""]
    contact = " | ".join(md_escape(p) for p in content.header.contact_parts())
    if contact:
        lines += [contact, ""]
    if content.summary:
        lines += ["## Summary", "", md_escape(content.summary), ""]
    if content.education:
        lines += ["## Education", ""]
        for e in content.education:
            degree, inst, dates = education_lines(e)
            lines.append(f"**{md_escape(degree)}**  ")
            lines.append(md_escape(inst) + ("  " if dates else ""))
            if dates:
                lines.append(md_escape(dates))
            lines += [f"- {md_escape(b.text)}" for b in e.bullets]
            lines.append("")
    for heading, entries in (("Experience", content.experience), ("Projects", content.projects)):
        if not entries:
            continue
        lines += [f"## {heading}", ""]
        for x in entries:
            title, detail = entry_lines(x)
            lines.append(f"**{md_escape(title)}**" + ("  " if detail else ""))
            if detail:
                lines.append(md_escape(detail))
            lines += [f"- {md_escape(b.text)}" for b in x.bullets]
            lines.append("")
    if content.skills:
        lines += ["## Skills", "", md_escape(", ".join(content.skills)), ""]
    if content.certifications:
        lines += ["## Certifications", ""] + [f"- {md_escape(c)}" for c in content.certifications] + [""]
    return "\n".join(lines).rstrip() + "\n"


def cover_letter_md(content: CoverLetterContent) -> str:
    lines = [f"# {md_escape(content.header.full_name)}", ""]
    contact = " | ".join(md_escape(p) for p in content.header.contact_parts())
    if contact:
        lines += [contact, ""]
    lines += [md_escape(content.date_text), "", md_escape(content.greeting), ""]
    for p in content.paragraphs:
        lines += [md_escape(p), ""]
    lines += [md_escape(content.closing) + "  ", md_escape(content.signature)]
    return "\n".join(lines) + "\n"


def render_resume(content: ResumeContent, fmt: str) -> bytes:
    if fmt == "docx":
        return resume_docx(content)
    if fmt == "pdf":
        return resume_pdf(content)
    if fmt == "md":
        return resume_md(content).encode("utf-8")
    raise PermanentError(f"unsupported document format {fmt!r}")


def render_cover_letter(content: CoverLetterContent, fmt: str) -> bytes:
    if fmt == "docx":
        return cover_letter_docx(content)
    if fmt == "pdf":
        return cover_letter_pdf(content)
    if fmt == "md":
        return cover_letter_md(content).encode("utf-8")
    raise PermanentError(f"unsupported document format {fmt!r}")
